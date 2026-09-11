"""Build auditable descriptive statistics and standalone comparison artifacts."""
import csv, hashlib, html, json, math, zipfile
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from classify import AREAS

BASE=Path(__file__).resolve().parent
OUT=BASE/'output'; OUT.mkdir(exist_ok=True)
papers=json.loads((BASE/'papers_classified_internal.json').read_text())
catalog_path=BASE.parent/'output/catalog.json'
catalog=json.loads(catalog_path.read_text())
sources=json.loads((BASE/'sources.json').read_text())
excluded=json.loads((BASE/'excluded.json').read_text())
CZ={
'q':'Kvantové počítání','graph':'Grafy a grafové algoritmy','learn':'Teorie učení','crypto':'Kryptografie',
'approx':'Aproximační algoritmy','ds':'Datové struktury a komprese','alg':'Algebraické a numerické výpočty',
'code':'Kódování a teorie informace','game':'Algoritmická teorie her a dělení zdrojů','online':'Online algoritmy, bandité, stochastická optimalizace',
'cx':'Výpočetní a obvodová složitost','count':'Enumerace a počítání','proof':'Důkazová složitost a logika',
'opt':'Optimalizace a matematické programování','comb':'Kombinatorika a grafové polynomy','dist':'Distribuované a lokální algoritmy',
'stream':'Streaming a sketching','prg':'Pseudonáhodnost a derandomizace','avg':'Average-case složitost','dyn':'Dynamické grafové algoritmy',
'geom':'Výpočetní geometrie','test':'Property testing a učení distribucí','csp':'Constraint satisfaction',
'param':'Parametrizované a exaktní algoritmy','fg':'Fine-grained složitost','comm':'Komunikační složitost',
'dp':'Diferenciální soukromí','lattice':'Mřížky a výpočetní teorie čísel','general':'Obecný návrh algoritmů',
'sched':'Rozvrhování a packing','comp':'Vyčíslitelnost a algoritmická informace','inf':'Nekonečné stavové systémy a verifikace',
'verify':'Hry, syntéza a verifikace','temp':'Temporální grafové algoritmy','sat':'Automatické dokazování a unifikace',
'top':'Výpočetní topologie','autom':'Automaty a formální jazyky','bio':'Algoritmy pro biologické struktury',
'sem':'Přepisování, lambda kalkul a sémantika','db':'Databáze a konečná teorie modelů','etr':'Existenční teorie reálných čísel',
'iba':'Informační složitost numerických úloh','random':'Randomizované prohledávání a optimalizace'}
labels={AREAS[k]:v for k,v in CZ.items()}
ca=Counter(r['area'] for r in catalog)
assert set(ca)==set(AREAS.values()) and len(ca)==43
assert len(papers)==1546 and len(catalog)==4695
assert len({r['doi'] for r in papers})==len(papers)
assert all(r['area'] in ca and r['doi'] and r['abstract'] for r in papers)

def comparison(subset):
    counts={v:Counter(r['area'] for r in subset if r['venue']==v) for v in ['STOC','FOCS']}
    sizes=Counter(r['venue'] for r in subset)
    rows=[]
    for area,n in ca.items():
        s,f=counts['STOC'][area],counts['FOCS'][area]
        cp=n/len(catalog); pp=(s+f)/len(subset)
        rows.append(dict(area=area,area_cs=labels[area],stoc_count=s,focs_count=f,papers_count=s+f,
          stoc_pct=100*s/sizes['STOC'],focs_pct=100*f/sizes['FOCS'],papers_pct=100*pp,
          catalog_count=n,catalog_pct=100*cp,catalog_minus_papers_pp=100*(cp-pp),
          catalog_to_papers_share_ratio=cp/pp if pp else None,
          expected_catalog_count_at_paper_shares=pp*len(catalog)))
    return sorted(rows,key=lambda x:(-x['papers_count'],-x['catalog_count'],x['area']))

rows=comparison(papers)
common=comparison([r for r in papers if 2022<=r['year']<=2025])
year_rows=[]
for src in sources:
    year_rows.append({k:src[k] for k in ['venue','year','count','url','metadata_url']})
area_year=[]
for src in sources:
    c=Counter(r['area'] for r in papers if r['venue']==src['venue'] and r['year']==src['year'])
    for a in ca:
        area_year.append(dict(venue=src['venue'],year=src['year'],area=a,count=c[a],pct=100*c[a]/src['count']))

public_fields=['id','venue','year','title','authors','area','url','doi','source_url','metadata_url','classification_method','classification_note','initial_area','rule_area']
public=[{k:r[k] for k in public_fields} for r in papers]
def write_csv(name,data):
    with (OUT/name).open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(data[0]));w.writeheader();w.writerows(data)
for name,data in [('comparison.csv',rows),('papers.csv',public),('editions.csv',year_rows),('areas_by_edition.csv',area_year),('common_years_2022_2025.csv',common)]:
    write_csv(name,data)
for name,data in [('comparison.json',rows),('papers.json',public),('sources.json',sources),('excluded.json',excluded)]:
    (OUT/name).write_text(json.dumps(data,ensure_ascii=False,indent=2))

notes=[
('Datum','2026-09-09'),
('Období','Posledních pět dokončených ročníků každé konference: STOC 2022–2026; FOCS 2021–2025. Jde o označení ročníků, ne shodné kalendářní okno.'),
('Jednotky','Konference: publikované výzkumné články. Katalog: záznamy indexu otevřených otázek. Srovnávají se podíly, nikoli absolutní množství problémů.'),
('Korpus','1 546 článků s DOI a abstraktem z oficiálních metadat ACM/IEEE. STOC 907; FOCS 639. Vyloučeny dvě keynote položky STOC 2024 a veškerá front/back matter.'),
('Katalog','4 695 záznamů ve 43 oblastech. Zachováno původní zařazení. Velká část katalogu vznikla automatickou extrakcí; nebyl zde znovu ověřován aktuální otevřený status ani přetříděn.'),
('Třídění','Vlastní orientační zařazení do jedné hlavní oblasti; nejde o oficiální taxonomii konferencí. Kandidátní pravidla doplnila kontrola názvů všech 1 546 článků modelem; u 43 vybraných nejednoznačných případů byla provedena kontrola abstraktu. Žádný nezávislý expertní audit.'),
('Srovnatelnost třídění','Názvy 43 kategorií jsou stejné jako v katalogu. Katalog a články byly tříděny odlišně; rozdíly proto obsahují i vliv metodiky, zejména u obecné zbytkové kategorie a mezi překrývajícími se obory.'),
('Překryvy','Každý článek se počítá jednou. Kvantová kryptografie se typicky řadí do kryptografie, kvantové kódy do kvantového počítání. Nula znamená žádný článek s tímto hlavním štítkem; nevylučuje přítomnost tématu.'),
('Konvence oblastí','Enumerace a počítání zahrnuje diskrétní sampling a mixing související s počítáním. Spojité log-concave sampling patří k optimalizaci. Důkazová složitost zahrnuje PCP/IOP a některé SoS lower bounds. Average-case zahrnuje planted modely a výpočetní prahy; statistické odhadování zpravidla teorii učení.'),
('Výpočet','Podíl oblasti = počet jejích článků / 1 546. Katalogový podíl = počet jejích záznamů / 4 695. Rozdíl je katalog minus články v procentních bodech. Ročníky se váží počtem publikovaných článků.'),
('Poměr','Poměr zastoupení = katalogový podíl / konferenční podíl. Pro nulu v konferenčním podílu je poměr nevyplněný. Očekávané počty při konferenčních podílech jsou pouze ilustrace jiné skladby, nikoli doporučená kvóta.'),
('Citlivost období','Soubor common_years_2022_2025.csv opakuje výpočet pro shodné ročníky 2022–2025 obou konferencí (1 218 článků).'),
('Interpretace','STOC/FOCS představují publikační profil dvou selektivních konferencí, nikoli reprezentativní vzorek celé TCS nebo míru významu oborů. Katalog čerpá i ze specializovaných konferencí, knih a sbírek.'),
('Kontrola zdrojů','FOCS 2025 accepted-papers-with-abstracts obsahoval starý seznam 2024 a nebyl použit. Čerpáno z vlastních IEEE proceedings metadat. FOCS 2023 má 142 publikovaných článků; předběžný seznam přijatých článků není jednotkou této statistiky.'),
('Stejné názvy','Dva články STOC 2022 s názvem Hypercontractivity on high dimensional expanders mají různé autory a DOI, proto jsou dvěma položkami.'),
('Reprodukce','Z pracovní složky postupně spustit parse.py, abstract_review.py, review.py, classify.py, build.py. Skripty používají původní cache; pro obnovení cache slouží fetch.py. Existující katalog se neupravuje.'),
]

wb=Workbook();wb.remove(wb.active)
for title,data in [('Oblasti',rows),('Rocniky',year_rows),('Clanky',public),('Oblasti_po_rocnicich',area_year),('Shodne_roky_2022_2025',common),('Metodika',[dict(tema=k,popis=v) for k,v in notes])]:
    ws=wb.create_sheet(title); ws.append(list(data[0]))
    for r in data: ws.append([r[k] for k in data[0]])
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]: cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='193F5C')
    for i,k in enumerate(data[0],1):
        ws.column_dimensions[get_column_letter(i)].width=65 if k in ['area','title','popis'] else 35 if 'url' in k or k in ['authors','area_cs'] else 20
        if 'pct' in k or k.endswith('_pp') or 'ratio' in k or 'expected_' in k:
            for cell in list(ws.columns)[i-1][1:]: cell.number_format='0.00'
    ws.row_dimensions[1].height=28
wb.save(OUT/'stoc-focs-vs-katalog.xlsx')

def fmt(x): return f'{x:.2f}'.replace('.',',')
def mdtable(data):
    t=['| Oblast | STOC | FOCS | Články % | Katalog: n (%) | Rozdíl p. b. |','|---|---:|---:|---:|---:|---:|']
    for r in data:
        t.append(f"| {r['area_cs']} | {r['stoc_count']} | {r['focs_count']} | {fmt(r['papers_pct'])} | {r['catalog_count']} ({fmt(r['catalog_pct'])}) | {r['catalog_minus_papers_pp']:+.2f} |".replace('.',','))
    return '\n'.join(t)
report=['# STOC / FOCS versus katalog otevřených otázek TCS','',
'**1 546 výzkumných článků proti 4 695 záznamům katalogu, 43 tematických oblastí.**','',
'## Ročníky a primární zdroje','', '| Konference | Ročník | Články | Zdroj |','|---|---:|---:|---|']
for s in sources: report.append(f"| {s['venue']} | {s['year']} | {s['count']} | [Oficiální sborník]({s['url']}) |")
report+=['','## Oblasti','','Rozdíl = podíl katalogu minus podíl konferenčních článků; kladně znamená vyšší zastoupení v katalogu.','',mdtable(rows),'','## Metodika a omezení','']
report += [f'- **{k}:** {v}' for k,v in notes]
(OUT/'report.md').write_text('\n'.join(report))
(OUT/'table.md').write_text(mdtable(rows))
(OUT/'README.md').write_text('\n\n'.join(f'**{k}**\n\n{v}' for k,v in notes))

# Standalone HTML: interactive comparison bars and searchable paper bibliography.
payload=json.dumps(dict(rows=rows,papers=public,sources=sources),ensure_ascii=False).replace('</','<\\/')
method=''.join(f'<p><strong>{html.escape(k)}:</strong> {html.escape(v)}</p>' for k,v in notes)
page='''<!doctype html><html lang="cs"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>STOC / FOCS × katalog TCS</title><style>
body{font:16px system-ui;max-width:1350px;margin:40px auto;padding:0 24px;color:#142f42;background:#f6f8fb}h1{font-size:34px;margin-bottom:8px}h2{margin-top:38px}a{color:#166ca5}p{line-height:1.55}table{border-collapse:collapse;width:100%;background:white}th,td{text-align:left;padding:9px;border-bottom:1px solid #e1e6ed}th{background:#193f5c;color:white;position:sticky;top:0}td.num{text-align:right;font-variant-numeric:tabular-nums}button,select,input{font:inherit;padding:9px;border:1px solid #bbc9d8;border-radius:5px;background:white;margin:4px}input{min-width:290px}.note{color:#506579}.cards{display:flex;gap:18px;flex-wrap:wrap}.card{padding:22px;background:white;border-radius:9px;flex:1}.card b{font-size:28px}.plotrow{display:grid;grid-template-columns:340px 1fr 160px;gap:14px;align-items:center;margin:10px 0;cursor:pointer}.track{height:12px;background:#e4e9ef;margin:3px 0;border-radius:3px}.bar{height:100%;border-radius:3px}.catalog{background:#d58e37}.papers{background:#277caa}.scroll{overflow:auto}.legend span{display:inline-block;padding:5px 12px;border-radius:4px;color:white}details{background:white;padding:18px;border-radius:8px;margin:24px 0}@media(max-width:760px){.plotrow{grid-template-columns:1fr}.card{min-width:130px}body{padding:0 12px}th,td{font-size:13px}input{min-width:200px}}
</style><h1>STOC / FOCS × katalog TCS</h1><p>Podíly oblastí v posledních pěti dokončených ročnících každé konference a v katalogu otevřených otázek. Stav k 9. 9. 2026.</p>
<div class="cards"><div class="card"><b>907</b><br>STOC 2022–2026</div><div class="card"><b>639</b><br>FOCS 2021–2025</div><div class="card"><b>4 695</b><br>záznamů katalogu</div><div class="card"><b>43</b><br>oblastí</div></div>
<p class="note">Jeden článek = jedna hlavní oblast. Vlastní orientační klasifikace podle názvů a vybraných abstraktů; publikační profil není měřítkem počtu otevřených problémů ani reprezentativní statistikou celé TCS.</p>
<p><a href="stoc-focs-vs-katalog.xlsx">Excel</a> · <a href="comparison.csv">Souhrn CSV</a> · <a href="papers.csv">Články CSV</a> · <a href="report.md">Zpráva Markdown</a></p>
<h2>Zastoupení oblastí</h2><div class="legend"><span class="papers">Konferenční články</span> <span class="catalog">Katalog</span></div>
<label>Konference <select id="venue"><option value="papers">STOC + FOCS</option value="stoc">STOC</option value="focs">FOCS</option></select></label>
<label>Řadit <select id="sort"><option value="gap">Velikost rozdílu podílů</option><option value="papers">Podíl konference</option><option value="catalog">Podíl katalogu</option></select></label>
<p class="note">Kliknutím na oblast se zobrazí její články. Všechny sloupce mají společné měřítko.</p><div id="plot"></div>
<h2>Úplná tabulka</h2><div class="scroll"><table id="summary"></table></div>
<h2 id="bibliography">Články a kontrola zařazení</h2><input id="query" placeholder="Hledat název, autora nebo DOI"><select id="area"><option value="">Všechny oblasti</option></select><select id="paperVenue"><option value="">Obě konference</option><option>STOC</option><option>FOCS</option></select><p id="matches"></p><div class="scroll"><table id="paperTable"></table></div>
<details><summary>Metodika, konvence a omezení</summary>METHOD</details><details><summary>Primární zdroje a počty ročníků</summary><div id="sources"></div></details>
<script>const D=PAYLOAD;
const E=s=>String(s??'').replace(/[&<>"']/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const pct=x=>x.toLocaleString('cs-CZ',{minimumFractionDigits:2,maximumFractionDigits:2});const $=id=>document.getElementById(id);
function plot(){let v=$('venue').value,k=v+'_pct';let rs=[...D.rows];let sort=$('sort').value;rs.sort((a,b)=>sort==='gap'?Math.abs(b.catalog_pct-b[k])-Math.abs(a.catalog_pct-a[k]):sort==='papers'?b[k]-a[k]:b.catalog_pct-a.catalog_pct);let max=Math.max(...rs.flatMap(r=>[r.catalog_pct,r[k]]));$('plot').innerHTML=rs.map(r=>`<div class="plotrow" data-area="${E(r.area)}"><span>${E(r.area_cs)}</span><div><div class="track"><div class="bar papers" style="width:${100*r[k]/max}%"></div></div><div class="track"><div class="bar catalog" style="width:${100*r.catalog_pct/max}%"></div></div></div><span>${pct(r[k])}% / ${pct(r.catalog_pct)}%</span></div>`).join('');document.querySelectorAll('.plotrow').forEach(e=>e.onclick=()=>{$('area').value=e.dataset.area;paperTable();$('bibliography').scrollIntoView({behavior:'smooth'})})}
$('summary').innerHTML='<tr><th>Oblast</th><th>STOC</th><th>FOCS</th><th>Články %</th><th>Katalog n</th><th>Katalog %</th><th>Rozdíl p. b.</th></tr>'+D.rows.map(r=>`<tr><td>${E(r.area_cs)}</td><td class="num">${r.stoc_count}</td><td class="num">${r.focs_count}</td><td class="num">${pct(r.papers_pct)}</td><td class="num">${r.catalog_count}</td><td class="num">${pct(r.catalog_pct)}</td><td class="num">${pct(r.catalog_minus_papers_pp)}</td></tr>`).join('');
$('area').innerHTML+=D.rows.map(r=>`<option value="${E(r.area)}">${E(r.area_cs)}</option>`).join('');
function paperTable(){let q=$('query').value.toLowerCase(),a=$('area').value,v=$('paperVenue').value;let rs=D.papers.filter(r=>(!a||r.area===a)&&(!v||r.venue===v)&&(!q||(r.title+' '+r.authors+' '+r.doi).toLowerCase().includes(q)));$('matches').textContent=rs.length+' článků'+(rs.length>150?' — zobrazeno prvních 150, celý seznam je v CSV a Excelu.':'');$('paperTable').innerHTML='<tr><th>Ročník</th><th>Článek</th><th>Oblast</th><th>Kontrola</th></tr>'+rs.slice(0,150).map(r=>`<tr><td>${r.venue} ${r.year}</td><td><a href="${E(r.url)}">${E(r.title)}</a><br><small>${E(r.authors)}</small></td><td>${E(r.area)}</td><td title="${E(r.classification_note)}">${r.classification_method.endsWith('abstract')?'Název + abstrakt':'Název'}</td></tr>`).join('')}
$('sources').innerHTML=D.sources.map(s=>`<p><a href="${s.url}">${s.venue} ${s.year}</a>: ${s.count} článků · <a href="${s.metadata_url}">metadata</a></p>`).join('');['venue','sort'].forEach(x=>$(x).onchange=plot);['area','paperVenue'].forEach(x=>$(x).onchange=paperTable);$('query').oninput=paperTable;plot();paperTable();
</script></html>'''.replace('METHOD',method).replace('PAYLOAD',payload)
(OUT/'comparison.html').write_text(page)

validation=dict(papers=len(papers),catalog=len(catalog),areas=len(rows),unique_dois=len({r['doi'] for r in papers}),
  stoc=sum(r['venue']=='STOC' for r in papers),focs=sum(r['venue']=='FOCS' for r in papers),
  classifications=Counter(r['classification_method'] for r in papers),excluded=excluded,
  count_totals_agree=sum(r['papers_count'] for r in rows)==len(papers) and sum(r['catalog_count'] for r in rows)==len(catalog),
  catalog_sha256=hashlib.sha256(catalog_path.read_bytes()).hexdigest(),
  source_cache_sha256={s['metadata_url']:hashlib.sha256((BASE.parents[1]/s['cache_file']).read_bytes()).hexdigest() for s in sources})
(OUT/'validation.json').write_text(json.dumps(validation,ensure_ascii=False,indent=2))
with zipfile.ZipFile(OUT/'stoc-focs-comparison.zip','w',zipfile.ZIP_DEFLATED) as z:
    for p in sorted(OUT.iterdir()):
        if p.suffix!='.zip':z.write(p,p.name)
    for p in BASE.glob('*.py'):z.write(p,'scripts/'+p.name)
    for f in ['overrides.json','abstract_review_ids.json']:z.write(BASE/f,'scripts/'+f)
print(json.dumps(validation,ensure_ascii=False,indent=2))
print('Top underrepresented',[(r['area'],round(r['catalog_minus_papers_pp'],2)) for r in sorted(rows,key=lambda r:r['catalog_minus_papers_pp'])[:8]])
print('Common-window top underrepresented',[(r['area'],round(r['catalog_minus_papers_pp'],2)) for r in sorted(common,key=lambda r:r['catalog_minus_papers_pp'])[:8]])
