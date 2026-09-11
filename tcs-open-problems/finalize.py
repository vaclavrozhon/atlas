from pathlib import Path
import json,re,csv,collections,hashlib,html,unicodedata
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parent;OUT=ROOT/'output';OUT.mkdir(exist_ok=True)
raw=[]
for f in ['base_rows','extra_rows','landmark_rows','dagstuhl_rows','supplement_rows']:raw+=json.loads((ROOT/(f+'.json')).read_text())
excluded=[]
for f in ['base_excluded','dagstuhl_excluded']:excluded+=json.loads((ROOT/(f+'.json')).read_text())
# Final review: exact formulations, source updates, subsequent primary results.
DROP={
('automata','20.8'):('The parity two-token conjecture was proved in 2025; broader extensions require their own formulations.','https://arxiv.org/abs/2503.24244'),
('sublinear','Problem 1'):('Old update-time target has extensive subsequent work; omitted pending a precise current formulation.',''),
('sublinear','Problem 2'):('Old mixed quantile-space and analysis questions have substantial later resolutions; omitted pending model-specific reformulation.',''),
('sublinear','Problem 11'):('The page’s update supplies the requested turnstile triangle-counting algorithm.',''),
('sublinear','Problem 12'):('Deterministic near-optimal CUR decompositions were subsequently constructed.','https://arxiv.org/abs/1405.7910'),
('sublinear','Problem 45'):('The page reports the linear-space Max-Cut lower bound; use the newer individually stated CSP conjectures.','https://www.cs.umd.edu/~gasarch/open/streamapprox.pdf'),
('sublinear','Problem 84'):('Polynomial-time approximate PML was subsequently obtained; the old mixed formulation is omitted.','https://arxiv.org/abs/1905.08448'),
('topp','Problem 67'):('Equal-area equal-perimeter convex partitions for arbitrary part counts were proved; the main conjecture is solved.','https://arxiv.org/abs/1804.03057'),
('rta','Problem 44'):('The source itself says that the core problem was solved (Boudet–Contejean).',''),
('rta','Problem 74'):('A source comment claims a resolution; omitted because the exact remaining target is unclear.',''),
('rta','Problem 88'):('The source comment explicitly reports a positive solution (Goubault-Larrecq, 1999).',''),
('Open Problems Related to Quantum Query Complexity','Problem 2'):('A quantum oracle separation was announced on 2026-09-02.','https://arxiv.org/abs/2609.02865'),
}
DK={
'21452:5.3':('The two-token conjecture was proved in 2025.','https://arxiv.org/abs/2503.24244'),
'11091:5.2':('The report says this was settled during the seminar.',''),
'22452:4.3':('The source labels this problem solved.',''),
'23351:4.3':('The report says the question was solved completely during the seminar.',''),
'13331:6.1':('The report’s update answers both questions affirmatively.',''),
'21171:4.3':('The report explicitly says these questions were later solved.',''),
'23162:4.9':('Upward Planarity is W[1]-hard parameterized by treewidth.','https://repositum.tuwien.at/handle/20.500.12708/230293'),
'13331:6.6':('Subsequent tight ETH lower bounds address this running-time question.',''),
'11081:5.4':('Duplicate of the more recent Automata Exchange separating-words entry.',''),
'13421:4.12':('Duplicate of the later planar subgraph-isomorphism difference-parameter question.',''),
}
# Do not treat partial-case resolutions as full resolutions in the audit.
for x in excluded:
 if x.get('source_url','').endswith('/56.html'):x['reason']='The countable case is solved in the source. The unrestricted uncountable formulation was omitted, not declared solved.'
 if 'one-pass cardinality' in x.get('reason',''):x['solution_url']='https://arxiv.org/abs/2607.14656'
# Explicitly record other screened questions even though they were omitted earlier.
excluded += [dict(title='Min-2-Lin over Z4 parameterized by deletions',source_url='https://pacs2024.github.io/pacs2024-open-problems.pdf',reason='The prime-power-modulus case was proved FPT in April 2026.',solution_url='https://arxiv.org/abs/2604.10369'),dict(title='Classical oracle separation of QMA and QCMA',source_url='https://www.scottaaronson.com/papers/open.pdf',reason='A classical oracle separation was announced in November 2025.',solution_url='https://arxiv.org/abs/2511.09551')]
RENAME={
'14071:4.2':'Complexity of editing a graph into one biclique and isolates',
'16241:4.5':'A trivariate hypertree polynomial reflecting triangulation triality',
'16221:4.1':'Near-linear algorithms for many-terminal vertex-disjoint surface paths',
'16221:4.3':'Weighted Max Cut parameterized by surface genus',
'16221:4.4':'Faster planar Steiner Tree algorithms parameterized by terminals',
'16221:4.7':'Planar Independent Set above the n/4 guarantee',
'16221:4.8':'Planar subgraph isomorphism parameterized by vertex or edge deficit',
'16221:4.10':'Near-linear-time minimum-cost Steiner perfect matching in planar graphs',
'15242:5.3':'Explicit dispersers for large quadratic-system solution sets',
'18361:5.7':'The reverse-mathematical strength of compact Hausdorff regularity',
'23331:5.2':'Complexity of balanced hypergraph partitioning and its variants',
'11071:4.7':'Optimal randomized ray-search strategies without order restrictions',
}
AREA_SEM={
'12021':'Computability and algorithmic information','18361':'Computability and algorithmic information','21461':'Computability and algorithmic information',
'14342':'Algorithmic game theory and fair division','25092':'Randomized search and optimization','25471':'Online and stochastic algorithms','25372':'Computational geometry','23351':'Information-based complexity and numerical algorithms','14071':'Parameterized and exact algorithms','18421':'Enumeration and counting','15242':'Algebraic and numerical computation'}
AREA_MAP={'Exact and exponential-time algorithms':'Parameterized and exact algorithms','Constraint satisfaction and parameterized complexity':'Constraint satisfaction','Online and stochastic algorithms':'Online algorithms, bandits and stochastic optimization','Online algorithms, bandits and reinforcement learning':'Online algorithms, bandits and stochastic optimization','Algorithms and complexity':'General algorithm design'}
COLLECTION={'automata':'Automata Exchange','sublinear':'Sublinear.info','topp':'The Open Problems Project','rta':'RTA Open Problems'}
rows=[]
for x in raw:
 x=dict(x);key=x.get('source_key','');why=DK.get(key) or DROP.get((x['source_collection'],x['source_locator']))
 if why:
  excluded.append(dict(title=x['title'],source_url=x['source_url'],reason=why[0],solution_url=why[1]));continue
 if key in RENAME:x['title']=RENAME[key]
 if x['source_collection']=='automata' and x['source_locator']=='19.5':x['title']='Complexity of universality for unambiguous context-free grammars'
 if x['source_collection']=='SIGACT Open Problems Column' and x['source_url'].endswith('/streamapprox.pdf'):
  if x['source_locator']=='Conjecture 3':x['title']='Linear-space lower bounds for two-pass Max-Cut approximation'
  if x['source_locator']=='Conjecture 9':x['title']='Beyond-square-root sketching space for near-half Max-DiCut approximation'
 sem=key.split(':')[0]
 if sem in AREA_SEM:x['area']=AREA_SEM[sem]
 x['area']=AREA_MAP.get(x['area'],x['area']);x['source_collection']=COLLECTION.get(x['source_collection'],x['source_collection'])
 x['authors']=re.sub(r'^By\s+','',x.get('authors','')).strip()
 # Preserve math notation but remove extraction control characters and HTML escaping.
 for k,v in list(x.items()):
  if isinstance(v,str):x[k]=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',html.unescape(v)).strip()
 if not x['title']:raise ValueError('empty title')
 # A source section containing several related questions remains explicitly marked as a family.
 if key in ['14071:4.3','14071:4.4','14071:4.6','19271:4.1','18421:4.1','18421:4.2','18421:4.8','18421:4.11','23331:5.2','16221:4.2','21171:4.5','12021:4.1','12021:4.2','12021:4.6','12021:4.7','18361:5.4','18361:5.5','25201:5.7','23121:4.1','19131:4.2']:x['scope']='problem_family'
 if x['source_collection']=='Sublinear.info' and x['source_locator'] in ['Problem 21','Problem 27','Problem 55','Problem 70','Problem 77','Problem 82','Problem 93']:x['scope']='problem_family'
 if x['source_collection']=='RTA Open Problems' and x['source_locator'] in ['Problem 26','Problem 91']:x['scope']='problem_family'
 if x['source_collection']=='The Open Problems Project' and x['source_locator']=='Problem 8':x['scope']='landmark'
 if x['status']=='maintainer_lists_open':x['source_year']=''
 x['year_basis']={'Automata Exchange':'contribution year','Sublinear.info':'workshop year listed on problem page','The Open Problems Project':'latest listed revision year','RTA Open Problems':'problem date listed on source page','Dagstuhl Reports':'seminar year; report may be published later','Jukka Suomela: locality problems':'last page update','Antoine Amarilli: research questions':'undated page','Landmark reference':'undated current listing'}.get(x['source_collection'],'source publication or workshop year')
 x['source_age_years']=2026-int(x['source_year']) if x['source_year'].isdigit() else ''
 x['reading_note']='Short index label; consult the source locator for the full formulation, definitions and assumptions.'
 rows.append(x)
# Exact title duplicates only: related mathematical variants are not silently collapsed.
seen={};kept=[]
for x in rows:
 key=re.sub(r'\W','',unicodedata.normalize('NFKC',x['title']).casefold())
 if key in seen:
  # Same title can denote different models; retain with a source context if so.
  if x['title']=='Wireless Scheduling':x['title']+=' ('+x['source_year']+' formulation)'
  else:
   old=seen[key];old['additional_source_urls']='; '.join(filter(None,[old.get('additional_source_urls',''),x['source_url']]))
   excluded.append(dict(title=x['title'],source_url=x['source_url'],reason='Duplicate title merged with the retained source entry.',solution_url=''));continue
 seen[key]=x;kept.append(x)
rows=kept
# Deterministic browsing order: landmarks first, then area and descending source year.
priority={t:i for i,t in enumerate(['P versus NP','NP versus coNP','P versus BPP','L versus NL','VP versus VNP','Unique Games Conjecture','Matrix multiplication exponent two','Linear Programming: Strongly Polynomial?'])}
rows.sort(key=lambda x:(priority.get(x['title'],100),x['scope']!='landmark',x['area'],-int(x['source_year'] or 0),x['title'].casefold()))
for i,x in enumerate(rows,1):x['id']=f'TCS-{i:04d}'
fields=['id','title','area','scope','source_year','year_basis','source_age_years','status','status_note','source_collection','source_locator','authors','source_url','pdf_url','additional_source_urls','source_title','accessed','reading_note']
rows=[{k:x.get(k,'') for k in fields} for x in rows]
# Audit repeats reflect the same omitted source, not extra problems.
seen=set();audit=[]
for x in excluded:
 k=(x['title'],x['source_url'],x['reason'])
 if k not in seen:seen.add(k);audit.append(x)
counts=collections.Counter(x['area'] for x in rows);years=collections.Counter(x['source_year'] or 'unknown' for x in rows)
sourcegroups=collections.defaultdict(list)
for x in rows:sourcegroups[x['source_url'].split('#')[0]].append(x)
sources=[]
for u,items in sorted(sourcegroups.items()):sources.append(dict(source_url=u,collection=items[0]['source_collection'],source_title=items[0]['source_title'],years=', '.join(sorted(set(x['source_year'] for x in items)-{''})),entry_count=len(items)))
for name,data in [('catalog',rows),('excluded',audit),('sources',sources)]:
 (OUT/(name+'.json')).write_text(json.dumps(data,ensure_ascii=False,indent=2))
 with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
  cols=list(data[0]);w=csv.DictWriter(f,cols,extrasaction='ignore');w.writeheader();w.writerows(data)
meta=dict(entry_count=len(rows),area_count=len(counts),distinct_source_documents=len(sources),scope_counts=dict(collections.Counter(x['scope'] for x in rows)),areas=dict(sorted(counts.items())),years=dict(sorted(years.items())),excluded_count=len(audit),compiled='2026-09-09',fully_reverified_current_status=False)
(OUT/'metadata.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2))
# Workbook.
wb=Workbook();intro=wb.active;intro.title='Start';intro.append(['TCS: katalog otevřených výzkumných otázek',''])
intro_rows=[('Počet',len(rows)),('Oblasti',len(counts)),('Zdroje',len(sources)),('Sestaveno','2026-09-09'),('Co katalog obsahuje','Stručné názvy konkrétních otázek s primárním zdrojem a lokátorem. Úplné definice a předpoklady jsou ve zdroji.'),('Aktuálnost','Většina položek: zdroj je označuje jako otevřené; současný stav nebyl individuálně a vyčerpávajícím způsobem ověřen. Nejde o seznam garantovaně nevyřešený k dnešku.'),('Rozsah položky','focused = konkrétní zadání; problem_family = zdrojem vymezená skupina otázek; landmark = významná obecná otázka. Nejde o odhad obtížnosti.'),('Rok','Rok zdroje nebo semináře, nikoli nutně rok prvního položení problému. Prázdný znamená, že rok není bezpečně doložen.'),('Kontrola','Odstraněny zachycené vyřešené položky, duplicity a neurčité agendy. Úplnou sémantickou deduplikaci ani plošnou rešerši všech pozdějších řešení nelze z tohoto katalogu vyvozovat.'),('Použití','Na listu Catalog filtrujte oblast, rok a scope. Otevřete source_url a ve zdroji najděte source_locator. Před volbou výzkumného projektu ověřte navazující literaturu.'),('List Excluded','Audit vynechaných kandidátů. Zahrnuje řešení, duplicity i nejisté nebo nevhodně vymezené položky; ne všechny jsou vyřešené.')]
for r in intro_rows:intro.append(r)
intro.column_dimensions['A'].width=26;intro.column_dimensions['B'].width=115
for row in intro:
 for c in row:c.alignment=Alignment(vertical='top',wrap_text=True)
for i in range(1,intro.max_row+1):intro.row_dimensions[i].height=48 if i>5 else 26

def sheet(name,data,columns):
 ws=wb.create_sheet(name);ws.append(columns)
 for row in data:ws.append([row.get(k,'') for k in columns])
 ws.freeze_panes='C2';ws.auto_filter.ref=ws.dimensions
 for c in ws[1]:c.fill=PatternFill('solid',fgColor='152F46');c.font=Font(color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
 ws.row_dimensions[1].height=32
 widths={'id':14,'title':65,'area':36,'scope':20,'source_year':14,'source_age_years':14,'status':27,'status_note':64,'source_collection':35,'source_locator':47,'authors':35,'source_url':58,'pdf_url':58,'additional_source_urls':58,'source_title':48,'accessed':14,'reading_note':60,'reason':85,'solution_url':65,'collection':35,'years':15,'entry_count':15,'count':15}
 for j,k in enumerate(columns,1):ws.column_dimensions[ws.cell(1,j).column_letter].width=widths.get(k,36)
 for row in ws.iter_rows(min_row=2):
  for c in row:
   c.alignment=Alignment(vertical='top',wrap_text=True)
   if isinstance(c.value,str) and c.value.startswith('http') and '; ' not in c.value:c.hyperlink=c.value;c.font=Font(color='126B87',underline='single')
  ws.row_dimensions[row[0].row].height=47
 if data:
  tab=Table(displayName=name+'Table',ref=ws.dimensions);tab.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);ws.add_table(tab)
 return ws
sheet('Catalog',rows,fields)
sheet('Areas',[{'area':k,'count':v} for k,v in sorted(counts.items(),key=lambda a:-a[1])],['area','count'])
sheet('Sources',sources,list(sources[0]))
sheet('Excluded',audit,['title','source_url','reason','solution_url'])
wb.save(OUT/'tcs-open-problems.xlsx')
assert load_workbook(OUT/'tcs-open-problems.xlsx',read_only=True)['Catalog'].max_row==len(rows)+1
print(json.dumps(meta,ensure_ascii=False,indent=2))
