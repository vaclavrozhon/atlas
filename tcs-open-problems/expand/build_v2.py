"""Publish the expanded research index while preserving every v1 identifier."""
from pathlib import Path
import collections,csv,hashlib,html,json,re,shutil,unicodedata,zipfile
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
from taxonomy import classify

ROOT=Path(__file__).resolve().parent;BASE=ROOT.parent;OUT=BASE/'output'
DATE='2026-09-09'
REVIEW_LABELS={'inherited_v1':'Původní katalog','manual_source_formulation':'Ručně zpracované zadání','automatic_source_passage':'Automatický výběr ze zdroje'}
AUTHORS={'codeequiv2026':'Anna-Lena Horlemann; Abhinaba Mazumder; Michael Schaller; Violetta Weger','eccc2023_019':'Pooya Hatami; William M. Hoza','eccc2023_094':'Lijie Chen; Roei Tell','eccc2025_083':'C. S. Bhargav; Prateek Dwivedi; Nitin Saxena','eccc2026_045':'Edward Pyne; Roei Tell','etr2024':'Marcus Schaefer; Jean Cardinal; Tillmann Miltzow','fair2022':'Georgios Amanatidis; Georgios Birmpas; Aris Filos-Ratsikas; Alexandros A. Voudouris','mixedfair2023':'Shengxin Liu; Xinhang Lu; Mashbat Suzuki; Toby Walsh'}

def norm(s):return re.sub(r'\W','',unicodedata.normalize('NFKC',s).casefold())
def clean(s):return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',str(s)).strip()
def get(fn):return json.loads((ROOT/fn).read_text())
def dump(name,rows):
    (OUT/(name+'.json')).write_text(json.dumps(rows,ensure_ascii=False,indent=2))
    if isinstance(rows,list) and rows:
        fields=list(dict.fromkeys(k for r in rows for k in r))
        with (OUT/(name+'.csv')).open('w',encoding='utf-8-sig',newline='') as f:
            w=csv.DictWriter(f,fields);w.writeheader()
            for row in rows:w.writerow({k:json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v for k,v in row.items()})

def main():
    old=json.loads((BASE/'output-v1-1003/catalog.json').read_text());rows=[]
    for x in old:
        x=dict(x);x.update(review_level='inherited_v1',statement_form='index_label',source_kind='original_collection',question_excerpt='',excerpt_truncated=False,candidate_id='',doi='',source_page='',selection_method='v1 retained',extraction_flags='',classification_basis='v1 classification')
        if x['id'] in ['TCS-0068','TCS-0069','TCS-0070']:x['area']='Computational and circuit complexity'
        if x['id']=='TCS-0071':x['area']='Proof complexity and logic'
        rows.append(x)
    manual={x['candidate_id']:x for x in get('manual_candidates.json')}
    source_info={x['key']:x for x in get('manual_sources.json')}
    for c in csv.DictReader((ROOT/'curated.tsv').open(),delimiter='\t'):
        x=manual[c['candidate_id']];info=source_info[x['manual_key']];key=x['manual_key'];year=x['year']
        basis='year of the cited source version; not the date of first posing'
        if key=='mixedfair2023':year=2024;basis='arXiv v4, 12 August 2024; first version 2023'
        if key=='raoyehudayoff':year='';basis='undated author manuscript; 2020 is the book publication year, not verified for this PDF'
        if key=='jukna2012':basis='2012 book publication year; cited author PDF is marked draft and is not independently dated'
        rows.append(dict(id='',title=c['title'],area=c['area'],scope='focused',source_year=str(year),year_basis=basis,source_age_years=2026-int(year) if year else '',status='source_open_unverified',status_note='The cited source poses this question; later literature was not exhaustively reviewed.',source_collection=info['title'],source_locator=x['source_locator']+f'; PDF page {x["page"]}',authors=AUTHORS.get(key,x.get('authors') or info.get('authors') or ''),source_url=x['source_url'],pdf_url=x['pdf_url'],additional_source_urls='',source_title=x['paper_title'],accessed=DATE,reading_note='Manually selected and paraphrased index label. Exact parameters, definitions and quantifiers are in the cited source.',review_level='manual_source_formulation',statement_form='manual_paraphrase',source_kind=info['kind'],question_excerpt='',excerpt_truncated=False,candidate_id=x['candidate_id'],doi=x['doi'],source_page=x['page'],selection_method='manual selection from numbered or explicitly open source statements',extraction_flags='',classification_basis='manual'))
    auto=get('selected_internal.json')
    # Strong lexical near-duplicates are merged, keeping the more recent source.
    # This cannot establish mathematical equivalence; weaker overlaps remain visible.
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.neighbors import NearestNeighbors
    auto.sort(key=lambda x:(-x['year'],-x['selection_score'],x['doi']))
    overlaps=[];removed=set()
    if auto:
        mat=TfidfVectorizer(analyzer='char_wb',ngram_range=(4,5),min_df=1).fit_transform(x['selected_sentence'] for x in auto)
        nn=NearestNeighbors(metric='cosine',n_neighbors=min(4,len(auto)),n_jobs=2).fit(mat)
        dist,idx=nn.kneighbors(mat)
        for i in range(len(auto)):
            for d,j in zip(dist[i],idx[i]):
                j=int(j);sim=1-float(d)
                if j>=i or sim<.65:continue
                action='possible lexical overlap; mathematical equivalence not assessed'
                if sim>=.94 and j not in removed:
                    removed.add(i);action='near-identical sentence merged into newer record'
                overlaps.append(dict(candidate_id=auto[i]['candidate_id'],other_candidate_id=auto[j]['candidate_id'],source_url=auto[i]['source_url'],other_source_url=auto[j]['source_url'],similarity=round(sim,4),action=action))
    accepted=[]
    for i,x in enumerate(auto):
        if i in removed:continue
        sentence=clean(x['selected_sentence']);words=sentence.split();truncated=len(words)>25
        excerpt=' '.join(words[:25])+(' […]' if truncated else '')
        author=x.get('authors') or '';author='; '.join(author) if isinstance(author,list) else author
        accepted.append(x)
        rows.append(dict(id='',title=excerpt,area=classify(x),scope='source_question_passage',source_year=str(x['year']),year_basis='proceedings publication year; later versions and solutions not systematically checked',source_age_years=2026-int(x['year']),status='source_passage_current_status_unverified',status_note='Automatic extraction found an unresolved-question statement or a question in a concluding section. Interpretation and present open status need individual review.',source_collection=x['venue'],source_locator=x['source_locator']+f'; PDF page {x["page"]}',authors=author,source_url=x['source_url'],pdf_url=x['pdf_url'],additional_source_urls='',source_title=x['paper_title'],accessed=DATE,reading_note='Short source excerpt, at most 25 words. Read it together with the article title and the linked PDF page; omitted text, notation and assumptions are not reconstructed. This row is a research lead, not a verified formal statement.',review_level='automatic_source_passage',statement_form='source_excerpt',source_kind='conference_paper',question_excerpt=excerpt,excerpt_truncated=truncated,candidate_id=x['candidate_id'],doi=x['doi'],source_page=x['page'],selection_method=x['extraction_kind'],extraction_flags='; '.join(x['flags']),classification_basis='heuristic from question, article title, publisher CCS and venue'))
    # Do not force the requested approximate size by inventing or splitting questions.
    for i,x in enumerate(rows):
        if not x['id']:x['id']=f'TCS-{i+1:04d}'
        for k,v in x.items():
            if isinstance(v,str):x[k]=clean(v)
    fields=list(rows[0]);rows=[{k:r.get(k,'') for k in fields} for r in rows]
    groups=collections.defaultdict(list)
    for x in rows:groups[x['source_url'].split('#')[0]].append(x)
    sources=[dict(source_url=u,source_title=a[0]['source_title'],collection=a[0]['source_collection'],source_kind=a[0]['source_kind'],authors=a[0]['authors'],years=', '.join(sorted({x['source_year'] for x in a}-{''})),entry_count=len(a)) for u,a in sorted(groups.items())]
    areas=collections.Counter(x['area'] for x in rows);reviews=collections.Counter(x['review_level'] for x in rows)
    meta_files=['lipics_metadata.json','extra_metadata.json','older_metadata.json','early_metadata.json']
    volumes={v['volume_url']:v for fn in meta_files if (ROOT/fn).exists() for v in get(fn)}
    pmlr=[x for fn in ['pmlr_metadata.json','pmlr_older_metadata.json'] if (ROOT/fn).exists() for x in get(fn)]
    meta=dict(entry_count=len(rows),new_entry_count=len(rows)-len(old),preserved_v1_entries=len(old),area_count=len(areas),distinct_source_documents=len(sources),review_counts=dict(reviews),scope_counts=dict(collections.Counter(x['scope'] for x in rows)),source_kind_counts=dict(collections.Counter(x['source_kind'] for x in rows)),areas=dict(sorted(areas.items(),key=lambda x:-x[1])),years=dict(sorted(collections.Counter(x['source_year'] or 'unknown' for x in rows).items())),compiled=DATE,fully_reverified_current_status=False,fully_semantically_deduplicated=False,conference_volumes_searched=len(volumes)+len({x['volume_url'] for x in pmlr}),conference_paper_files_searched=len(list((ROOT/'papers').glob('*.json'))),manual_sources_searched=len(source_info),manual_sources_retained=len({x['source_url'] for x in rows if x['review_level']=='manual_source_formulation'}),automatic_source_papers=len(accepted),near_identical_sentences_merged=len(removed),conference_series_retained=sorted({x['source_collection'] for x in rows if x['source_kind']=='conference_paper'}),method_version='2.0')
    excluded=json.loads((BASE/'output-v1-1003/excluded.json').read_text())+get('status_updates.json')
    dump('catalog',rows);dump('sources',sources);dump('excluded',excluded);dump('metadata',meta);dump('possible_overlaps',overlaps)
    dump('screening_audit',get('screening_audit.json'))
    dump('areas',[dict(area=k,count=n) for k,n in areas.most_common()])
    write_readme(meta);write_markdown(rows,areas);write_workbook(rows,sources,excluded,meta,areas,fields)
    payload=json.dumps(rows,ensure_ascii=False).replace('</','<\\/')
    template=(ROOT/'viewer.html').read_text().replace('__ROWS__',payload).replace('__META__',json.dumps(meta,ensure_ascii=False))
    (OUT/'catalog.html').write_text(template)
    # Reproducibility metadata includes consulted sources without redistributing their texts.
    consulted=[{k:v for k,v in x.items() if k not in ['pages','text','abstract','papers','complete']} for x in source_info.values()]
    dump('consulted_books_surveys',consulted)
    dump('consulted_conferences',[dict(volume_url=x['volume_url'],volume_title=x['volume_title'],year=x['year'],venue=x['venue'],paper_count=len(x['papers'])) for x in volumes.values()]+[dict(volume_url=u,volume_title=a[0]['volume_title'],year=a[0]['year'],venue=a[0]['venue'],paper_count=len(a)) for u,a in group_pmlr(pmlr).items()])
    validation=validate(rows,old,accepted,meta)
    dump('validation',validation)
    with zipfile.ZipFile(OUT/'tcs-open-problems.zip','w',zipfile.ZIP_DEFLATED,compresslevel=8) as z:
        for f in sorted(OUT.iterdir()):
            if f.is_file() and f.suffix in ['.md','.csv','.json','.html','.xlsx']:z.write(f,f.name)
    (OUT/'build-stats.txt').write_text(json.dumps(meta,ensure_ascii=False,indent=2))
    print(json.dumps(meta,ensure_ascii=False,indent=2))

def group_pmlr(pmlr):
    out=collections.defaultdict(list)
    for x in pmlr:out[x['volume_url']].append(x)
    return out

def write_readme(m):
    text=f'''# TCS: rozšířený katalog výzkumných otázek

Sestaveno {DATE}. **{m['entry_count']:,} záznamů v {m['area_count']} oblastech**, {m['distinct_source_documents']:,} zdrojových dokumentů. Přibylo {m['new_entry_count']:,} záznamů. ID původních 1 003 položek zůstala zachována.

## Co přesně znamená počet

- **{m['review_counts']['inherited_v1']:,}** záznamů z původního katalogu.
- **{m['review_counts']['manual_source_formulation']:,}** nových ručně vybraných a přeformulovaných zadání z knih, monografií a survey.
- **{m['review_counts']['automatic_source_passage']:,}** nových automaticky vybraných otázek ze stejných počtů různých konferenčních článků. Každá má krátký úryvek, název článku a odkaz na stránku PDF.

**Aktuální otevřenost většiny záznamů nebyla individuálně ověřena.** Automatická část je rešeršní index kandidátů: filtr hledal číslované otevřené otázky, explicitně nevyřešená tvrzení a otázky v závěrech. I po filtraci mohou zůstat chyby interpretace, neúplná matematická sazba, překryvy a otázky vyřešené po vydání zdroje. Počet záznamů proto není certifikovaný počet nezávislých problémů otevřených k {DATE}.

Původních 30 rodin problémů zůstává označeno `problem_family`. Automatická položka označuje jednu vybranou zdrojovou pasáž; pasáž může formulovat více souvisejících variant. Obtížnost menších úloh nebyla odhadována. `focused` neznamená snadný problém.

## Jak katalog používat

Otevřete **[catalog.html](catalog.html)**: funguje lokálně bez serveru, umožňuje hledání, filtrování podle oblasti, roku, způsobu zpracování a typu zdroje a export výběru. U automatického záznamu čtěte společně úryvek a název článku. `[…]` značí zkrácení; plná otázka, definice a kvantifikátory jsou na uvedené stránce zdroje.

- [Excel](tcs-open-problems.xlsx): listy Start, Catalog, Areas, Sources a Excluded.
- [CSV](catalog.csv) a [JSON](catalog.json): úplná strukturovaná data.
- [Markdown](catalog.md): celý seznam rozdělený podle oblastí.
- [Zdroje](sources.csv), [prohledané knihy a survey](consulted_books_surveys.csv), [prohledané sborníky](consulted_conferences.csv).
- [Vyřazené položky](excluded.csv), [automatický filtr](screening_audit.csv), [možné textové překryvy](possible_overlaps.csv), [validace](validation.json).

## Rozsah rešerše a výběr

Prohledáno {m['conference_volumes_searched']} konferenčních svazků a {m['conference_paper_files_searched']:,} dostupných textů článků, dále {m['manual_sources_searched']} knih, monografií, survey a sbírek. Ručně zpracované doplnění čerpá ze {m['manual_sources_retained']} z těchto zdrojů; ne každý prohledaný zdroj přinesl použitelný záznam. Výběr zvýhodňuje veřejně dostupné zdroje LIPIcs, PMLR, ECCC a autorské kopie. Není to úplná ani rovnoměrná bibliografie TCS; sborníky STOC, FOCS, SODA a kryptografických konferencí nebyly plošně vytěženy.

Zachyceny konferenční řady: {', '.join(m['conference_series_retained'])}.

Z jednoho konferenčního článku se automaticky přidává nejvýše jedna položka, bez rozdělování parametrů na umělé varianty. Krátké citace mají nejvýše 25 slov na článek. Automaticky se odstraňují zjevné důkazy, řešené otázky, pouhé odkazy na otázky, obecné agendy a přesná opakování; {m['near_identical_sentences_merged']} téměř shodných formulací bylo sloučeno. Matematická deduplikace napříč odlišnými formulacemi není dokončena. `possible_overlaps.csv` zaznamenává zbývající silnější textové podobnosti.

## Evidence a data

`review_level` rozlišuje převzatý index, ruční práci se zadáním a automatický výběr. Žádná z těchto hodnot sama nepotvrzuje dnešní otevřenost. `status` a `status_note` tuto nejistotu zachovávají. Rok znamená rok citované verze nebo sborníku, nikoli nutně rok položení otázky. U nedatovaného rukopisu zůstává prázdný.

`source_locator` uvádí číslo otázky nebo umístění pasáže a stránku PDF; PDF stránka se může lišit od tištěného číslování. `selection_method` a `extraction_flags` umožňují audit automatického výběru. Zařazení automatických položek do oblastí je heuristické; článek může zasahovat do více oblastí.

Při dílčí kontrole aktualizací byly vyřazeny mimo jiné některé staré otázky o citlivosti booleovských funkcí, faktorizaci formulí, property testingu a férovém rozdělování. Audit rozlišuje vyřešené položky od duplicit a nedostatečně určitých kandidátů. Nová správná bibliografická adresa ještě neznamená ověření současného stavu matematické otázky.

Zdrojové PDF a úplné vytěžené texty jsou v místní pracovní cache, nejsou součástí exportního ZIP. Skripty v sousedním adresáři `expand/` zachovávají postup sběru, výběru, ruční rozhodnutí a tvorbu výstupu. Původní katalog je v `output-v1-1003/`.
'''
    (OUT/'README.md').write_text(text)

def write_markdown(rows,areas):
    lines=[f'# TCS: {len(rows):,} záznamů výzkumných otázek','',f'Sestaveno {DATE}. Úrovně zpracování: původní index / ruční doplnění / automatický výběr. Současná otevřenost většiny položek není ověřena. Automatické úryvky je nutné číst s názvem článku a plným zdrojem. Podrobnosti v [README](README.md).','']
    for area,n in areas.most_common():
        lines += [f'## {area} ({n})','']
        for x in rows:
            if x['area']!=area:continue
            title=x['title'].replace('[','\\[').replace(']','\\]')
            lines.append(f'- **{x["id"]} — {title}** ({x["source_year"] or "nedatováno"}; {REVIEW_LABELS[x["review_level"]]}). [{x["source_title"] or x["source_collection"]}]({x["pdf_url"] or x["source_url"]}) — {x["source_locator"]}.')
        lines.append('')
    (OUT/'catalog.md').write_text('\n'.join(lines))

def write_workbook(rows,sources,excluded,m,areas,fields):
    wb=Workbook();ws=wb.active;ws.title='Start'
    for row in [('TCS – katalog výzkumných otázek',len(rows)),('Sestaveno',DATE),('Původní katalog',m['preserved_v1_entries']),('Ručně zpracované nové otázky',m['review_counts']['manual_source_formulation']),('Automatický výběr z článků',m['automatic_source_papers']),('Aktuálnost','Většina současných stavů nebyla individuálně ověřena. Automatická část obsahuje kandidáty pro další rešerši; nejde o certifikovaný seznam nezávislých současných otevřených problémů.'),('Čtení','U automatických položek čtěte úryvek společně s názvem článku a stránkou zdroje. […] znamená zkrácení. Matematické definice a kvantifikátory jsou v PDF.'),('Filtry','review_level, area, source_year a source_kind. Původní ID zůstala zachována.'),('Metodika','Viz README.md v exportu; zdroje a vyřazené otázky mají samostatné listy.')]:ws.append(row)
    ws.column_dimensions['A'].width=36;ws.column_dimensions['B'].width=110
    for row in ws:
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=45
    def sheet(name,data,cols):
        s=wb.create_sheet(name);s.append(cols);s.freeze_panes='C2'
        for r in data:
            s.append([r.get(k,'') for k in cols])
        for j,k in enumerate(cols,1):s.column_dimensions[s.cell(1,j).column_letter].width={'id':14,'title':70,'area':36,'source_title':65,'source_locator':45,'source_url':55,'pdf_url':60}.get(k,28)
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='183B50')
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment=Alignment(wrap_text=True,vertical='top')
                if isinstance(c.value,str):
                    c.data_type='s'
                    if c.value.startswith('http') and '; ' not in c.value:c.hyperlink=c.value;c.font=Font(color='0A6C8B',underline='single')
            s.row_dimensions[row[0].row].height=54
        if data:
            t=Table(displayName=name+'Table',ref=s.dimensions);t.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True);s.add_table(t)
    sheet('Catalog',rows,fields);sheet('Areas',[dict(area=k,count=n) for k,n in areas.most_common()],['area','count']);sheet('Sources',sources,list(sources[0]));sheet('Excluded',excluded,['title','source_url','reason','solution_url'])
    wb.save(OUT/'tcs-open-problems.xlsx')

def validate(rows,old,accepted,m):
    assert len(rows)==len({x['id'] for x in rows})
    assert [x['id'] for x in rows[:len(old)]]==[x['id'] for x in old]
    assert len(accepted)==len({x['doi'] for x in accepted})
    assert all(x['title'] and x['area'] and x['source_url'].startswith(('https://','http://')) for x in rows)
    assert all(len(x['question_excerpt'].replace(' […]','').split())<=25 for x in rows if x['review_level']=='automatic_source_passage')
    for x in accepted:
        r=json.loads(Path(x['source_file']).read_text())
        assert 1<=x['page']<=len(r['pages'])
        assert x['doi']==r['doi']
    with (OUT/'catalog.csv').open(encoding='utf-8-sig') as f:assert len(list(csv.DictReader(f)))==len(rows)
    wb=load_workbook(OUT/'tcs-open-problems.xlsx',read_only=True);assert wb['Catalog'].max_row==len(rows)+1;wb.close()
    return dict(passed=True,entry_count=len(rows),unique_ids=True,v1_ids_preserved=True,one_automatic_record_per_paper=True,excerpts_at_most_25_words_per_paper=True,automatic_pdf_page_locators_in_range=True,csv_json_xlsx_counts_agree=True,current_open_status_fully_verified=False,mathematical_deduplication_complete=False,checked=DATE)

if __name__=='__main__':main()
